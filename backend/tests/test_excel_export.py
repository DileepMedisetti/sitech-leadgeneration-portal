import io
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import MagicMock
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.db.database import get_db
from app.main import app
from app.models.lead import Lead
from app.services.excel_export_service import (
    generate_leads_excel_bytes,
    generate_search_results_excel_bytes,
    sanitize_cell_value,
)

client = TestClient(app)


def test_sanitize_cell_value_formula_injection():
    # Should prepend single quote for formula injection triggers
    assert sanitize_cell_value("=cmd|'/C calc'!A0") == "'=cmd|'/C calc'!A0"
    assert sanitize_cell_value("+1234567890") == "'+1234567890"
    assert sanitize_cell_value("-1+1") == "'-1+1"
    assert sanitize_cell_value("@SUM(A1:A10)") == "'@SUM(A1:A10)"
    assert sanitize_cell_value("\tmalicious") == "'\tmalicious"

    # Regular values should remain unchanged
    assert sanitize_cell_value("Dr. Jane Doe") == "Dr. Jane Doe"
    assert sanitize_cell_value("jane.doe@stanford.edu") == "jane.doe@stanford.edu"
    assert sanitize_cell_value(None) == ""
    assert sanitize_cell_value(42) == 42
    assert sanitize_cell_value(0.95) == 0.95


def test_generate_leads_excel_bytes_structure():
    mock_lead_1 = SimpleNamespace(
        id=101,
        name="Dr. Alice Walker",
        email="alice.walker@mit.edu",
        profession="Associate Professor",
        university="Massachusetts Institute of Technology",
        department="Computer Science and AI",
        research_specialization="Reinforcement Learning",
        country="United States",
        linkedin_url="https://linkedin.com/in/alicewalker",
        phone_number="+1 617-253-1000",
        website="https://people.csail.mit.edu/awalker",
        source_url="https://csail.mit.edu/faculty/directory",
        qualification_status="qualified",
        confidence_score=0.92,
        qualification_reason="Verified CS faculty member with active academic profile",
        search_job_id=5,
        created_at=datetime(2026, 3, 15, 10, 30, 0, tzinfo=timezone.utc),
    )

    mock_lead_2 = SimpleNamespace(
        id=102,
        name="=HYPERLINK(\"http://attacker.com\",\"Click Here\")",  # formula injection attempt
        email="bob.smith@ox.ac.uk",
        profession="Full Professor",
        university="University of Oxford",
        department="Department of Engineering",
        research_specialization="Robotics",
        country="United Kingdom",
        linkedin_url="https://linkedin.com/in/bobsmith",
        phone_number=None,
        website=None,
        source_url="https://eng.ox.ac.uk/people",
        qualification_status="qualified",
        confidence_score=0.88,
        qualification_reason="Verified Engineering professor",
        search_job_id=5,
        created_at=datetime(2026, 3, 15, 11, 0, 0, tzinfo=timezone.utc),
    )

    excel_bytes = generate_leads_excel_bytes([mock_lead_1, mock_lead_2])
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000  # valid binary zip/xlsx file

    # Read back with openpyxl to verify valid xlsx structure
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Sitech Verified Leads" in wb.sheetnames
    ws = wb["Sitech Verified Leads"]

    # Verify header row contains the 8 required columns in order
    header_row = [cell.value for cell in ws[1]]
    assert header_row[:8] == [
        "Name",
        "Email",
        "Designation / Role",
        "Specialization / Profession",
        "University / Institution",
        "Country",
        "LinkedIn URL",
        "Institutional Webpage / Source URL",
    ]
    assert "Lead ID" in header_row

    # Verify row 2 data matching the 8 fields
    assert ws.cell(row=2, column=1).value == "Dr. Alice Walker"
    assert ws.cell(row=2, column=2).value == "alice.walker@mit.edu"
    assert ws.cell(row=2, column=3).value == "Associate Professor"
    assert ws.cell(row=2, column=4).value == "Reinforcement Learning"
    assert ws.cell(row=2, column=5).value == "Massachusetts Institute of Technology"
    assert ws.cell(row=2, column=6).value == "United States"
    assert ws.cell(row=2, column=7).value == "https://linkedin.com/in/alicewalker"
    assert ws.cell(row=2, column=8).value == "https://people.csail.mit.edu/awalker"
    assert ws.cell(row=2, column=14).value == "92%"

    # Verify formula injection is defused in row 3
    assert ws.cell(row=3, column=1).value.startswith("'=")


def test_export_leads_xlsx_endpoint(monkeypatch):
    mock_db = MagicMock()
    app.dependency_overrides[get_db] = lambda: mock_db

    fake_excel = b"PK\x03\x04test_excel_binary_content"

    monkeypatch.setattr(
        "app.routes.leads.export_leads_to_excel",
        lambda db, status, job_id: fake_excel,
    )

    response = client.get("/api/leads/export-xlsx?status=qualified")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "sitech_leads_qualified.xlsx" in response.headers["content-disposition"]
    assert response.content == fake_excel

    app.dependency_overrides.clear()


def test_generate_search_results_excel_bytes():
    mock_res_1 = SimpleNamespace(
        id=1,
        search_job_id=42,
        result_url="https://cs.stanford.edu/people",
        result_title="Faculty Directory | Stanford CS",
        status="processed",
        retry_count=0,
        error_message=None,
        lead_id=10,
        processed_at=datetime(2026, 3, 15, 12, 0, 0, tzinfo=timezone.utc),
        created_at=datetime(2026, 3, 15, 11, 30, 0, tzinfo=timezone.utc),
    )
    mock_res_2 = SimpleNamespace(
        id=2,
        search_job_id=42,
        result_url="https://mit.edu/faculty",
        result_title="=cmd|'/C calc'!A0",
        status="failed",
        retry_count=2,
        error_message="HTTP 403 Forbidden",
        lead_id=None,
        processed_at=datetime(2026, 3, 15, 12, 5, 0, tzinfo=timezone.utc),
        created_at=datetime(2026, 3, 15, 11, 31, 0, tzinfo=timezone.utc),
    )

    excel_bytes = generate_search_results_excel_bytes([mock_res_1, mock_res_2])
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Stored Extraction Results" in wb.sheetnames
    ws = wb["Stored Extraction Results"]

    headers = [cell.value for cell in ws[1]]
    assert "Result ID" in headers
    assert "Target Webpage URL" in headers
    assert "Webpage Title" in headers
    assert "Status" in headers
    assert "Extraction Details" in headers

    # Verify formula injection is defused
    assert ws.cell(row=3, column=4).value.startswith("'=")


def test_export_search_job_results_xlsx_endpoint(monkeypatch):
    mock_db = MagicMock()
    app.dependency_overrides[get_db] = lambda: mock_db

    fake_excel = b"PK\x03\x04test_search_results_binary"

    monkeypatch.setattr(
        "app.routes.search_jobs.get_job_detail",
        lambda db, job_id: {"id": job_id, "keywords": "ai professors"},
    )
    monkeypatch.setattr(
        "app.routes.search_jobs.export_search_results_to_excel",
        lambda db, job_id: fake_excel,
    )

    response = client.get("/api/search-jobs/42/results/export-xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "extraction_results_job_42.xlsx" in response.headers["content-disposition"]
    assert response.content == fake_excel

    app.dependency_overrides.clear()

