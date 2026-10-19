import io
import openpyxl
import pytest
from bs4 import BeautifulSoup
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import MagicMock

from app.services.extraction.cleaners import (
    clean_verification_url,
    extract_direct_profile_url,
    is_generic_homepage,
    is_valid_profile_url,
    select_lead_verification_url,
)
from app.services.lead_extraction_service import extract_all_contacts
from app.services.excel_export_service import (
    generate_leads_excel_bytes,
    get_lead_verification_url,
)
from app.services.lead_service import export_leads_to_csv
from app.repositories.lead_repository import enrich_existing_lead


# =========================================================================
# Case 1: Faculty profile extraction -> saves direct profile URL
# =========================================================================
def test_case_1_faculty_profile_extraction_saves_direct_profile_url():
    html = """
    <div class="faculty-profile">
        <h1 class="faculty-name">Dr. Johnathan Smith</h1>
        <div class="title">Professor of Computer Science</div>
        <div class="email">Email: <a href="mailto:jsmith@cs.stanford.edu">jsmith@cs.stanford.edu</a></div>
        <a class="bio-link" href="https://cs.stanford.edu/people/johndoe">Personal Profile & Lab</a>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://cs.stanford.edu/faculty/directory.html",
        search_criteria={"profession": "Computer Science"},
    )
    assert len(contacts) >= 1
    lead = contacts[0]
    assert lead["name"] == "Dr. Johnathan Smith"
    assert lead["email"] == "jsmith@cs.stanford.edu"
    assert lead["website"] == "https://cs.stanford.edu/people/johndoe"
    assert lead["source_url"] == "https://cs.stanford.edu/people/johndoe"
    assert not is_generic_homepage(lead["website"])


# =========================================================================
# Case 2: Laboratory page extraction -> saves lab/profile URL
# =========================================================================
def test_case_2_laboratory_page_extraction_saves_lab_or_profile_url():
    html = """
    <div class="lab-member-card">
        <h3>Dr. Alice Vance</h3>
        <p class="role">AI Researcher & Postdoctoral Scholar</p>
        <p><a href="mailto:avance@csail.mit.edu">avance@csail.mit.edu</a></p>
        <p><a href="/labs/vision-nlp/avance" class="lab-link">Vision & NLP Lab Page</a></p>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://csail.mit.edu/research/labs.html",
        search_criteria={"profession": "AI Researcher"},
    )
    assert len(contacts) >= 1
    lead = contacts[0]
    assert lead["name"] == "Dr. Alice Vance"
    assert lead["email"] == "avance@csail.mit.edu"
    # Specific lab profile URL must be resolved against domain
    assert lead["website"] == "https://csail.mit.edu/labs/vision-nlp/avance"
    assert lead["source_url"] == "https://csail.mit.edu/labs/vision-nlp/avance"


# =========================================================================
# Case 3: Department directory extraction -> saves directory entry URL
# =========================================================================
def test_case_3_department_directory_extraction_saves_entry_anchor_or_page():
    html = """
    <div class="directory-list">
        <div class="faculty-entry" id="prof-robert-chen">
            <h4>Dr. Robert Chen</h4>
            <span class="designation">Professor of Neurology</span>
            <p>Email: <a href="mailto:rchen@neurology.ucla.edu">rchen@neurology.ucla.edu</a></p>
        </div>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://neurology.ucla.edu/faculty/directory.html",
        search_criteria={"profession": "Neurology"},
    )
    assert len(contacts) >= 1
    lead = contacts[0]
    assert lead["name"] == "Dr. Robert Chen"
    # Container has id "prof-robert-chen", anchor is appended or directory page URL is preserved
    assert lead["website"] in (
        "https://neurology.ucla.edu/faculty/directory.html#prof-robert-chen",
        "https://neurology.ucla.edu/faculty/directory.html",
    )
    # Must never be domain root
    assert lead["website"] != "https://neurology.ucla.edu"
    assert lead["website"] != "https://ucla.edu"
    assert not is_generic_homepage(lead["website"])


# =========================================================================
# Case 4: Multiple candidates from same page -> each receives their own distinct URL
# =========================================================================
def test_case_4_multiple_candidates_from_same_page_distinct_urls():
    html = """
    <div class="directory-table">
        <div class="person-card">
            <h3>Dr. Clara Barton</h3>
            <div class="pos">Dentist</div>
            <a href="/faculty/cbarton">Profile of Dr. Barton</a>
            <a href="mailto:cbarton@dentistry.umich.edu">cbarton@dentistry.umich.edu</a>
        </div>
        <div class="person-card">
            <h3>Dr. Arthur Pendelton</h3>
            <div class="pos">Dental Surgeon</div>
            <a href="/faculty/apendelton">Profile of Dr. Pendelton</a>
            <a href="mailto:apendelton@dentistry.umich.edu">apendelton@dentistry.umich.edu</a>
        </div>
        <div class="person-card">
            <h3>Dr. Beatrice Lowe</h3>
            <div class="pos">Orthodontist</div>
            <a href="/faculty/blowe">Profile of Dr. Lowe</a>
            <a href="mailto:blowe@dentistry.umich.edu">blowe@dentistry.umich.edu</a>
        </div>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://dentistry.umich.edu/faculty-all.html",
        search_criteria={"profession": "Dentistry"},
    )
    assert len(contacts) == 3

    barton = next(c for c in contacts if "Barton" in c["name"])
    pendelton = next(c for c in contacts if "Pendelton" in c["name"])
    lowe = next(c for c in contacts if "Lowe" in c["name"])

    # Ensure URLs are distinct and uniquely mapped to each person
    assert barton["website"] == "https://dentistry.umich.edu/faculty/cbarton"
    assert pendelton["website"] == "https://dentistry.umich.edu/faculty/apendelton"
    assert lowe["website"] == "https://dentistry.umich.edu/faculty/blowe"

    assert len({barton["website"], pendelton["website"], lowe["website"]}) == 3


# =========================================================================
# Case 5: Generic homepage search result with direct profile available
# =========================================================================
def test_case_5_generic_homepage_prefers_direct_profile_or_rejects_generic():
    # Sub-case A: Homepage URL, but person has direct profile link in HTML
    html_with_profile = """
    <div>
        <h2>Dr. Sarah Jenkins</h2>
        <p>Professor of Gynecology</p>
        <a href="/departments/obgyn/faculty/sjenkins">Dr. Jenkins Faculty Page</a>
        <a href="mailto:sjenkins@harvard.edu">sjenkins@harvard.edu</a>
    </div>
    """
    contacts_a = extract_all_contacts(
        html=html_with_profile,
        page_url="https://harvard.edu/",
        search_criteria={"profession": "Gynecology"},
    )
    assert len(contacts_a) >= 1
    lead_a = contacts_a[0]
    # Direct profile must be selected, NOT generic homepage
    assert lead_a["website"] == "https://harvard.edu/departments/obgyn/faculty/sjenkins"
    assert not is_generic_homepage(lead_a["website"])

    # Sub-case B: Homepage URL, no direct profile link for person -> generic homepage rejected
    html_without_profile = """
    <div>
        <h2>Dr. Thomas Wayne</h2>
        <p>Professor of Gynecology</p>
        <a href="mailto:twayne@harvard.edu">twayne@harvard.edu</a>
    </div>
    """
    contacts_b = extract_all_contacts(
        html=html_without_profile,
        page_url="https://harvard.edu/",
        search_criteria={"profession": "Gynecology"},
    )
    assert len(contacts_b) >= 1
    lead_b = contacts_b[0]
    # Website and source_url must be None/empty, NOT https://harvard.edu/
    assert lead_b["website"] is None
    assert lead_b["source_url"] is None


# =========================================================================
# Case 6: No reliable source URL -> field left empty/None, never fabricated
# =========================================================================
def test_case_6_no_reliable_source_url_left_empty_never_fabricated():
    # Test clean_verification_url directly with search engine and generic URLs
    assert clean_verification_url("https://www.google.com/search?q=dr+smith") is None
    assert clean_verification_url("https://bing.com/search?q=dr+smith") is None
    assert clean_verification_url("https://stanford.edu/") is None
    assert clean_verification_url("https://ox.ac.uk/index.html") is None
    assert clean_verification_url(None) is None
    assert clean_verification_url("") is None

    # Test select_lead_verification_url fallback
    v_url, s_url = select_lead_verification_url(
        direct_url="https://google.com/url?q=...",
        page_url="https://berkeley.edu/",
        person_name="Dr. Jane Smith",
    )
    assert v_url is None
    assert s_url is None

    # Test repository enrich_existing_lead does not fabricate or overwrite with generic
    existing_lead = SimpleNamespace(
        id=42,
        name="Dr. Jane Smith",
        email="jsmith@berkeley.edu",
        profession="Associate Professor",
        research_specialization="Biology",
        department=None,
        university="UC Berkeley",
        country="United States",
        website="https://vcresearch.berkeley.edu/faculty/jsmith",
        source_url="https://vcresearch.berkeley.edu/faculty/jsmith",
        phone_number=None,
        linkedin_url=None,
        qualification_status="qualified",
        confidence_score=0.95,
        qualification_reason="Verified",
        evidence={},
        search_job_id=1,
    )

    # An incoming lead with a generic homepage must NOT overwrite existing specific profile
    incoming_candidate = {
        "website": "https://berkeley.edu/",
        "source_url": "https://berkeley.edu/",
        "evidence": {},
    }
    enrich_existing_lead(existing_lead, incoming_candidate)
    assert existing_lead.website == "https://vcresearch.berkeley.edu/faculty/jsmith"
    assert existing_lead.source_url == "https://vcresearch.berkeley.edu/faculty/jsmith"


# =========================================================================
# Case 7: Excel and CSV export test verifying Institutional Webpage / Source URL column
# =========================================================================
def test_case_7_excel_and_csv_export_institutional_webpage_source_url_column():
    lead_with_valid_url = SimpleNamespace(
        id=201,
        name="Dr. Maya Lin",
        email="mlin@yale.edu",
        profession="Professor of Dentistry",
        research_specialization="Dentistry",
        university="Yale University",
        country="United States",
        linkedin_url="https://linkedin.com/in/mayalin",
        website="https://dentistry.yale.edu/faculty/mlin",
        department="Dental Medicine",
        phone_number="+1 203-785-0000",
        source_url="https://dentistry.yale.edu/faculty/mlin",
        qualification_status="qualified",
        confidence_score=0.94,
        created_at=datetime(2026, 4, 15, 14, 0, 0, tzinfo=timezone.utc),
    )

    lead_with_generic_url = SimpleNamespace(
        id=202,
        name="Dr. Liam O'Connor",
        email="loconnor@tcd.ie",
        profession="Assistant Professor",
        research_specialization="Computer Science",
        university="Trinity College Dublin",
        country="Ireland",
        linkedin_url=None,
        website="https://www.tcd.ie/",
        department="School of Computer Science",
        phone_number=None,
        source_url="https://www.tcd.ie/index.html",
        qualification_status="qualified",
        confidence_score=0.88,
        created_at=datetime(2026, 4, 15, 14, 30, 0, tzinfo=timezone.utc),
    )

    leads = [lead_with_valid_url, lead_with_generic_url]

    # Verify Excel Export
    excel_bytes = generate_leads_excel_bytes(leads)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    assert "Institutional Webpage / Source URL" in headers
    url_col_idx = headers.index("Institutional Webpage / Source URL") + 1

    # Row 2 (lead_with_valid_url): should have the exact profile URL
    row_2_val = ws.cell(row=2, column=url_col_idx).value
    assert row_2_val == "https://dentistry.yale.edu/faculty/mlin"

    # Row 3 (lead_with_generic_url): generic homepage was filtered out, should be blank/empty
    row_3_val = ws.cell(row=3, column=url_col_idx).value
    assert row_3_val in ("", None)

    # Verify CSV Export
    csv_str = export_leads_to_csv(leads)
    lines = [line.strip() for line in csv_str.strip().split("\n")]
    assert len(lines) == 3

    csv_headers = lines[0].split(",")
    assert "Institutional Webpage / Source URL" in csv_headers
    csv_url_idx = csv_headers.index("Institutional Webpage / Source URL")

    csv_row_2 = lines[1].split(",")
    assert csv_row_2[csv_url_idx] == "https://dentistry.yale.edu/faculty/mlin"

    csv_row_3 = lines[2].split(",")
    assert csv_row_3[csv_url_idx] == ""
