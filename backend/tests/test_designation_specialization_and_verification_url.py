import io
import openpyxl
import pytest
from bs4 import BeautifulSoup
from types import SimpleNamespace
from datetime import datetime, timezone

from app.services.extraction.cleaners import (
    check_profession_match,
    extract_designation_and_specialization,
    extract_direct_profile_url,
    is_valid_profile_url,
)
from app.services.extraction.qualifier import calculate_qualification_score
from app.services.lead_extraction_service import extract_all_contacts, extract_contact_details
from app.services.excel_export_service import generate_leads_excel_bytes
from app.services.lead_service import export_leads_to_csv


# =========================================================================
# 1. DESIGNATION & SPECIALIZATION EXTRACTION ACROSS DISCIPLINES
# =========================================================================

def test_dentistry_extraction_designation_and_specialization():
    html = """
    <div class="faculty-directory">
        <div class="card" id="dr-smith">
            <h3>Dr. Michael Smith</h3>
            <div class="role">Professor of Dentistry</div>
            <div class="department">Department of Pediatric Dentistry</div>
            <a href="/faculty/msmith" class="profile-link">View Profile</a>
            <a href="mailto:msmith@nyu.dental.edu">msmith@nyu.dental.edu</a>
        </div>
        <div class="card" id="dr-jones">
            <h3>Dr. Sarah Jones</h3>
            <div class="role">Dental Surgeon</div>
            <a href="/people/sjones">Dr. Sarah Jones Bio</a>
            <a href="mailto:sjones@nyu.dental.edu">sjones@nyu.dental.edu</a>
        </div>
        <div class="card" id="dr-miller">
            <h3>Dr. David Miller</h3>
            <div class="role">Orthodontist</div>
            <a href="mailto:dmiller@nyu.dental.edu">dmiller@nyu.dental.edu</a>
        </div>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://dental.nyu.edu/faculty/directory.html",
        search_criteria={"profession": "Dentistry"},
    )
    assert len(contacts) == 3

    smith = next(c for c in contacts if "Smith" in c["name"])
    assert "Professor of Dentistry" in smith["profession"]
    assert smith["research_specialization"] == "Dentistry"
    assert smith["website"] == "https://dental.nyu.edu/faculty/msmith"
    assert smith["qualification_status"] == "qualified"

    jones = next(c for c in contacts if "Jones" in c["name"])
    assert jones["profession"] == "Dental Surgeon"
    assert jones["research_specialization"] == "Dentistry"
    assert jones["website"] == "https://dental.nyu.edu/people/sjones"
    assert jones["qualification_status"] == "qualified"

    miller = next(c for c in contacts if "Miller" in c["name"])
    assert miller["profession"] == "Orthodontist"
    assert miller["research_specialization"] == "Dentistry"
    # No direct anchor link, so falls back to container anchor or full page_url
    assert miller["website"] in ("https://dental.nyu.edu/faculty/directory.html#dr-miller", "https://dental.nyu.edu/faculty/directory.html")
    assert miller["qualification_status"] == "qualified"


def test_neurology_extraction_designation_and_specialization():
    html = """
    <div class="faculty-list">
        <div class="member">
            <h4>Dr. Elena Rostova</h4>
            <span class="designation">Professor of Neurology</span>
            <div class="specialization">Specialization: Clinical Neurophysiology and Epilepsy</div>
            <a href="https://med.stanford.edu/profiles/elena-rostova">View Faculty Profile</a>
            <a href="mailto:erostova@stanford.edu">erostova@stanford.edu</a>
        </div>
        <div class="member">
            <h4>Dr. Alan Turing</h4>
            <span class="designation">Clinical Neurologist</span>
            <a href="/labs/brain-center/aturing">Lab Page</a>
            <a href="mailto:aturing@stanford.edu">aturing@stanford.edu</a>
        </div>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://med.stanford.edu/neurology/people.html",
        search_criteria={"profession": "Neurology"},
    )
    assert len(contacts) == 2

    elena = next(c for c in contacts if "Elena" in c["name"])
    assert "Professor of Neurology" in elena["profession"]
    assert "Clinical Neurophysiology" in elena["research_specialization"]
    assert elena["website"] == "https://med.stanford.edu/profiles/elena-rostova"
    assert elena["qualification_status"] == "qualified"

    alan = next(c for c in contacts if "Alan" in c["name"])
    assert alan["profession"] == "Clinical Neurologist"
    assert alan["research_specialization"] == "Neurology"
    assert alan["website"] == "https://med.stanford.edu/labs/brain-center/aturing"
    assert alan["qualification_status"] == "qualified"


def test_gynecology_extraction_designation_and_specialization():
    html = """
    <table>
        <tr>
            <td>Dr. Mary Patterson</td>
            <td>Professor of Gynecology</td>
            <td><a href="/faculty/mpatterson">Profile</a></td>
            <td><a href="mailto:mpatterson@jhmi.edu">mpatterson@jhmi.edu</a></td>
        </tr>
        <tr>
            <td>Dr. Priya Sharma</td>
            <td>Obstetrician and Gynecologist</td>
            <td><a href="/staff/psharma">Bio</a></td>
            <td><a href="mailto:psharma@jhmi.edu">psharma@jhmi.edu</a></td>
        </tr>
    </table>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://www.hopkinsmedicine.org/gynecology/faculty",
        search_criteria={"profession": "Gynecology"},
    )
    assert len(contacts) == 2

    mary = next(c for c in contacts if "Mary" in c["name"])
    assert "Professor of Gynecology" in mary["profession"]
    assert mary["research_specialization"] == "Gynecology"
    assert mary["website"] == "https://www.hopkinsmedicine.org/faculty/mpatterson"
    assert mary["qualification_status"] == "qualified"

    priya = next(c for c in contacts if "Priya" in c["name"])
    assert "Obstetrician" in priya["profession"]
    assert priya["research_specialization"] == "Gynecology"
    assert priya["website"] == "https://www.hopkinsmedicine.org/staff/psharma"
    assert priya["qualification_status"] == "qualified"


def test_computer_science_extraction_designation_and_specialization():
    html = """
    <div class="team-grid">
        <div class="person-card">
            <h3>Dr. Geoffrey Hinton</h3>
            <p class="role">Computer Science Professor</p>
            <p class="research">Area of Research: Deep Learning and Neural Computation</p>
            <a href="https://cs.toronto.edu/~hinton">Homepage</a>
            <a href="mailto:hinton@cs.toronto.edu">hinton@cs.toronto.edu</a>
        </div>
        <div class="person-card">
            <h3>Dr. Fei-Fei Li</h3>
            <p class="role">Associate Professor of Computer Science</p>
            <a href="/people/feifeili">Full Bio</a>
            <a href="mailto:feifeili@stanford.edu">feifeili@stanford.edu</a>
        </div>
        <div class="person-card">
            <h3>Dr. Yann LeCun</h3>
            <p class="role">AI Researcher</p>
            <a href="/faculty/yann">Faculty Page</a>
            <a href="mailto:yann@nyu.edu">yann@nyu.edu</a>
        </div>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://cs.stanford.edu/directory",
        search_criteria={"profession": "Computer Science"},
    )
    assert len(contacts) == 3

    hinton = next(c for c in contacts if "Hinton" in c["name"])
    assert hinton["profession"] == "Computer Science Professor"
    assert "Deep Learning" in hinton["research_specialization"]
    assert hinton["website"] == "https://cs.toronto.edu/~hinton"

    feifei = next(c for c in contacts if "Fei-Fei" in c["name"])
    assert "Associate Professor of Computer Science" in feifei["profession"]
    assert feifei["research_specialization"] == "Computer Science"
    assert feifei["website"] == "https://cs.stanford.edu/people/feifeili"

    yann = next(c for c in contacts if "Yann" in c["name"])
    assert yann["profession"] == "AI Researcher"
    assert yann["research_specialization"] == "Computer Science"
    assert yann["website"] == "https://cs.stanford.edu/faculty/yann"


def test_engineering_extraction_designation_and_specialization():
    html = """
    <div class="directory">
        <div class="faculty-box">
            <h3>Dr. Robert Taylor</h3>
            <div class="title">Professor of Mechanical Engineering</div>
            <div class="interest">Specialization: Fluid Dynamics and Thermal Sciences</div>
            <a href="/faculty/rtaylor">Faculty Profile</a>
            <a href="mailto:rtaylor@mit.edu">rtaylor@mit.edu</a>
        </div>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://meche.mit.edu/faculty",
        search_criteria={"profession": "Mechanical Engineering"},
    )
    assert len(contacts) == 1
    lead = contacts[0]
    assert lead["profession"] == "Professor of Mechanical Engineering"
    assert "Fluid Dynamics" in lead["research_specialization"]
    assert lead["website"] == "https://meche.mit.edu/faculty/rtaylor"
    assert lead["qualification_status"] == "qualified"


# =========================================================================
# 2. PROFESSION RELEVANCE MATCHING & REJECTION OF UNQUALIFIED CANDIDATES
# =========================================================================

def test_profession_match_and_rejection_logic():
    # 1. Matching candidate
    match, err = check_profession_match(
        extracted_profession="Professor of Dentistry",
        extracted_specialization="Dentistry",
        requested_profession="Dentistry",
    )
    assert match is True
    assert err is None

    # 2. Non-matching candidate: requested Dentistry, candidate is Civil Engineering
    match, err = check_profession_match(
        extracted_profession="Professor of Civil Engineering",
        extracted_specialization="Civil Engineering",
        requested_profession="Dentistry",
    )
    assert match is False
    assert "does not match requested profession 'Dentistry'" in err

    # 3. Non-matching candidate: requested Neurology, candidate is Consultant Gynecologist
    match, err = check_profession_match(
        extracted_profession="Consultant Gynecologist",
        extracted_specialization="Gynecology",
        requested_profession="Neurology",
    )
    assert match is False
    assert "Gynecology" in err

    # 4. Qualifier rejects non-matching candidates
    status, score, reason, evidence = calculate_qualification_score(
        name="Dr. Charles Brown",
        email="cbrown@univ.edu",
        strategy="faculty_card",
        association_strength="card_container",
        profession="Professor of Mechanical Engineering",
        research_specialization="Mechanical Engineering",
        requested_profession="Dentistry",
    )
    assert status == "rejected"
    assert score == 0.0
    assert evidence["rejection_reason"] == "profession_mismatch"
    assert "does not match requested profession 'Dentistry'" in reason


def test_mixed_directory_rejects_unrelated_professions():
    html = """
    <div class="multi-department-directory">
        <div class="card">
            <h3>Dr. Alice Dentist</h3>
            <p class="role">Associate Professor of Dentistry</p>
            <a href="mailto:adentist@univ.edu">adentist@univ.edu</a>
        </div>
        <div class="card">
            <h3>Dr. Bob Engineer</h3>
            <p class="role">Professor of Mechanical Engineering</p>
            <a href="mailto:bengineer@univ.edu">bengineer@univ.edu</a>
        </div>
    </div>
    """
    contacts = extract_all_contacts(
        html=html,
        page_url="https://univ.edu/faculty-directory",
        search_criteria={"profession": "Dentistry"},
    )
    assert len(contacts) == 2

    dentist = next(c for c in contacts if "Alice" in c["name"])
    assert dentist["qualification_status"] == "qualified"

    engineer = next(c for c in contacts if "Bob" in c["name"])
    assert engineer["qualification_status"] == "rejected"
    assert engineer["evidence"]["rejection_reason"] == "profession_mismatch"


# =========================================================================
# 3. DIRECT INSTITUTIONAL PROFILE / LAB URL EXTRACTION PRIORITY
# =========================================================================

def test_direct_profile_url_priority_and_filtering():
    # Priority 1: Anchor text matching person name
    html1 = BeautifulSoup(
        '<div><a href="/faculty/jsmith">Dr. John Smith</a><a href="/dept/faculty">All Faculty</a></div>',
        "html.parser",
    )
    url1 = extract_direct_profile_url(html1, "https://harvard.edu/people", person_name="Dr. John Smith")
    assert url1 == "https://harvard.edu/faculty/jsmith"

    # Priority 2: Profile keywords ('View Profile')
    html2 = BeautifulSoup(
        '<div><h3>Dr. Alice</h3><a href="/profiles/alice-walker">View Profile</a><a href="/">Home</a></div>',
        "html.parser",
    )
    url2 = extract_direct_profile_url(html2, "https://harvard.edu/dept", person_name="Dr. Alice")
    assert url2 == "https://harvard.edu/profiles/alice-walker"

    # Priority 3: Path containing /people/ or /faculty/
    html3 = BeautifulSoup(
        '<div><a href="/people/bobby">Bobby Link</a></div>',
        "html.parser",
    )
    url3 = extract_direct_profile_url(html3, "https://harvard.edu/directory")
    assert url3 == "https://harvard.edu/people/bobby"

    # Disallowed URLs (Google search, LinkedIn, file downloads) must be rejected
    html_bad = BeautifulSoup(
        '<div><a href="https://www.google.com/search?q=dr+smith">Search</a><a href="/bio.pdf">CV</a></div>',
        "html.parser",
    )
    url_bad = extract_direct_profile_url(html_bad, "https://harvard.edu/directory/page.html")
    # Falls back to the page_url, NEVER root netloc or google
    assert url_bad == "https://harvard.edu/directory/page.html"
    assert url_bad != "https://harvard.edu"
    assert "google" not in url_bad

    # Fallback to element anchor ID if no link is present
    html_id = BeautifulSoup(
        '<div id="prof-johnson"><h3>Dr. Johnson</h3></div>',
        "html.parser",
    )
    url_id = extract_direct_profile_url(html_id.div, "https://harvard.edu/directory.html")
    assert url_id == "https://harvard.edu/directory.html#prof-johnson"


def test_is_valid_profile_url():
    assert is_valid_profile_url("https://med.stanford.edu/profiles/jane") is True
    assert is_valid_profile_url("https://cs.ox.ac.uk/people/alan") is True
    assert is_valid_profile_url("https://www.google.com/search?q=test") is False
    assert is_valid_profile_url("https://linkedin.com/in/test") is False
    assert is_valid_profile_url("https://univ.edu/paper.pdf") is False
    assert is_valid_profile_url("mailto:test@univ.edu") is False
    assert is_valid_profile_url(None) is False


# =========================================================================
# 4. EXCEL AND CSV EXPORT: 8 CRITICAL FIELDS
# =========================================================================

def test_excel_export_includes_eight_critical_fields():
    lead = SimpleNamespace(
        id=42,
        name="Dr. Catherine Wu",
        email="cwu@columbia.edu",
        profession="Associate Professor of Neurology",
        research_specialization="Cognitive Neurology",
        university="Columbia University",
        country="United States",
        linkedin_url="https://linkedin.com/in/catherinewu",
        website="https://neurology.columbia.edu/profile/cwu",
        department="Department of Neurology",
        phone_number="+1 212-305-2500",
        source_url="https://neurology.columbia.edu/faculty",
        qualification_status="qualified",
        confidence_score=0.95,
        qualification_reason="Verified neurology faculty member",
        search_job_id=7,
        created_at=datetime(2026, 4, 1, 10, 0, 0, tzinfo=timezone.utc),
    )

    excel_bytes = generate_leads_excel_bytes([lead])
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    expected_eight = [
        "Name",
        "Email",
        "Designation / Role",
        "Specialization / Profession",
        "University / Institution",
        "Country",
        "LinkedIn URL",
        "Institutional Webpage / Lab URL",
    ]
    assert headers[:8] == expected_eight

    # Verify Row 2 values
    row_values = [cell.value for cell in ws[2]]
    assert row_values[0] == "Dr. Catherine Wu"
    assert row_values[1] == "cwu@columbia.edu"
    assert row_values[2] == "Associate Professor of Neurology"
    assert row_values[3] == "Cognitive Neurology"
    assert row_values[4] == "Columbia University"
    assert row_values[5] == "United States"
    assert row_values[6] == "https://linkedin.com/in/catherinewu"
    assert row_values[7] == "https://neurology.columbia.edu/profile/cwu"


def test_csv_export_includes_eight_critical_fields():
    lead = SimpleNamespace(
        id=88,
        name="Dr. Rajesh Patel",
        email="rpatel@aiims.edu",
        profession="Professor of Dentistry",
        research_specialization="Dentistry",
        university="AIIMS New Delhi",
        country="India",
        linkedin_url="https://linkedin.com/in/rajeshpatel",
        website="https://aiims.edu/faculty/rpatel",
        department="Dental Sciences",
        phone_number="+91 11 2658 8500",
        source_url="https://aiims.edu/dental-faculty",
        qualification_status="qualified",
        confidence_score=0.91,
        created_at=datetime(2026, 4, 1, 12, 0, 0, tzinfo=timezone.utc),
    )

    csv_data = export_leads_to_csv(leads=[lead])
    lines = [line.strip() for line in csv_data.strip().split("\n")]
    assert len(lines) == 2

    header_cols = lines[0].split(",")
    # Verify critical fields exist in CSV header
    for expected_field in [
        "Person Name",
        "Professional Email",
        "Country",
        "University",
        "LinkedIn Profile URL",
        "Academic Designation",
        "Specialization / Profession",
        "Institutional Webpage / Lab URL",
    ]:
        assert expected_field in header_cols

    data_cols = lines[1].split(",")
    name_idx = header_cols.index("Person Name")
    email_idx = header_cols.index("Professional Email")
    desig_idx = header_cols.index("Academic Designation")
    spec_idx = header_cols.index("Specialization / Profession")
    uni_idx = header_cols.index("University")
    country_idx = header_cols.index("Country")
    li_idx = header_cols.index("LinkedIn Profile URL")
    web_idx = header_cols.index("Institutional Webpage / Lab URL")

    assert data_cols[name_idx] == "Dr. Rajesh Patel"
    assert data_cols[email_idx] == "rpatel@aiims.edu"
    assert data_cols[desig_idx] == "Professor of Dentistry"
    assert data_cols[spec_idx] == "Dentistry"
    assert data_cols[uni_idx] == "AIIMS New Delhi"
    assert data_cols[country_idx] == "India"
    assert data_cols[li_idx] == "https://linkedin.com/in/rajeshpatel"
    assert data_cols[web_idx] == "https://aiims.edu/faculty/rpatel"
