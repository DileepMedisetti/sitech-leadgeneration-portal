import io
import uuid
from datetime import datetime, timezone
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.db.database import SessionLocal
from app.main import app
from app.models.campaign import Campaign, CampaignRecipientSnapshot
from app.models.lead import Lead
from app.services.campaign_tracking_service import classify_bounce


def _utc_now():
    return datetime.now(timezone.utc)


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def db():
    session = SessionLocal()
    yield session
    session.close()


def test_classify_bounce_categories():
    # Hard bounces
    assert classify_bounce("550 5.1.1 User unknown") == "hard"
    assert classify_bounce("554 delivery error: dd Sorry your message to invalid@domain was rejected") == "hard"
    assert classify_bounce("Mailbox unavailable or disabled") == "hard"
    assert classify_bounce("Recipient address rejected: domain not found") == "hard"

    # Soft bounces
    assert classify_bounce("452 4.2.2 Mailbox full / quota exceeded") == "soft"
    assert classify_bounce("451 4.7.1 Temporarily deferred, try again later") == "soft"
    assert classify_bounce("Server busy: connection timed out") == "soft"

    # Unknown
    assert classify_bounce("Unspecified remote transmission warning") == "unknown"
    assert classify_bounce(None) == "unknown"


def test_completed_campaigns_overview_and_filtering(client: TestClient, db: Session):
    uid = uuid.uuid4().hex[:8]
    now = _utc_now()
    c1 = Campaign(
        name=f"Bioinformatics Outreach {uid}",
        status="sent",
        sent_at=now,
        total_recipients=10,
        sent_count=10,
        delivered_count=9,
        opened_count=5,
        clicked_count=2,
        bounced_count=1,
    )
    c2 = Campaign(
        name=f"Genomics Conference {uid}",
        status="completed",
        sent_at=now,
        total_recipients=5,
        sent_count=5,
        delivered_count=5,
        opened_count=3,
        clicked_count=1,
        bounced_count=0,
    )
    c_draft = Campaign(
        name=f"Draft Unsent Campaign {uid}",
        status="draft",
        total_recipients=2,
    )
    db.add_all([c1, c2, c_draft])
    db.commit()

    # Query completed campaigns
    res = client.get("/api/campaign-tracking/campaigns")
    assert res.status_code == 200
    data = res.json()

    # Completed campaigns list should exclude draft
    names = [item["name"] for item in data["items"]]
    assert f"Bioinformatics Outreach {uid}" in names
    assert f"Genomics Conference {uid}" in names
    assert f"Draft Unsent Campaign {uid}" not in names

    # Overview metrics
    metrics = data["overview_metrics"]
    assert metrics["completed_campaigns"] >= 2
    assert metrics["total_sent"] >= 15
    assert metrics["total_delivered"] >= 14
    assert metrics["total_bounces"] >= 1


def test_campaign_tracking_detail_and_opened_recipients(client: TestClient, db: Session):
    uid = uuid.uuid4().hex[:8]
    now = _utc_now()
    campaign = Campaign(
        name=f"Microbiology Symposium {uid}",
        status="sent",
        sent_at=now,
        total_recipients=3,
        sent_count=3,
        delivered_count=2,
        opened_count=2,
        clicked_count=1,
        bounced_count=1,
    )
    db.add(campaign)
    db.flush()

    # Add leads and snapshots
    lead1 = Lead(name="Dr. Alice Smith", email=f"alice_{uid}@oxford.ac.uk", profession="Professor", university="Oxford", country="United Kingdom")
    lead2 = Lead(name="Dr. Bob Jones", email=f"bob_{uid}@harvard.edu", profession="Researcher", university="Harvard", country="United States")
    lead3 = Lead(name="Dr. Carol White", email=f"carol_{uid}@invalid.edu", profession="Lecturer", university="Cambridge", country="United Kingdom")
    db.add_all([lead1, lead2, lead3])
    db.flush()

    # Snapshots
    s1 = CampaignRecipientSnapshot(
        campaign_id=campaign.id,
        lead_id=lead1.id,
        name=lead1.name,
        normalized_email=lead1.email,
        status="sent",
        dispatch_attempt_at=now,
        delivered_at=now,
        opened_at=now,
        open_count=2,
        clicked_at=now,
        click_count=1,
    )
    s2 = CampaignRecipientSnapshot(
        campaign_id=campaign.id,
        lead_id=lead2.id,
        name=lead2.name,
        normalized_email=lead2.email,
        status="sent",
        dispatch_attempt_at=now,
        delivered_at=now,
        opened_at=now,
        open_count=1,
    )
    s3 = CampaignRecipientSnapshot(
        campaign_id=campaign.id,
        lead_id=lead3.id,
        name=lead3.name,
        normalized_email=lead3.email,
        status="bounced",
        dispatch_attempt_at=now,
        bounced_at=now,
        bounce_reason="550 5.1.1 User unknown",
        bounce_type="hard",
    )
    db.add_all([s1, s2, s3])
    db.commit()

    # Test summary endpoint
    res_sum = client.get(f"/api/campaign-tracking/campaigns/{campaign.id}/summary")
    assert res_sum.status_code == 200
    sum_data = res_sum.json()
    assert sum_data["metrics"]["unique_opens"] == 2
    assert sum_data["metrics"]["historical_bounced_count"] == 1
    assert sum_data["metrics"]["hard_bounces"] == 1

    # Test opened recipients endpoint
    res_opened = client.get(f"/api/campaign-tracking/campaigns/{campaign.id}/opened-recipients")
    assert res_opened.status_code == 200
    opened_data = res_opened.json()
    assert opened_data["total"] == 2
    opened_emails = [item["email"] for item in opened_data["items"]]
    assert f"alice_{uid}@oxford.ac.uk" in opened_emails
    assert f"bob_{uid}@harvard.edu" in opened_emails
    assert f"carol_{uid}@invalid.edu" not in opened_emails
    assert "United Kingdom" in opened_data["available_countries"]

    # Test Excel export for opened recipients
    res_excel = client.get(f"/api/campaign-tracking/campaigns/{campaign.id}/opened-recipients/export")
    assert res_excel.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res_excel.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content))
    ws = wb.active
    assert ws.title == "Opened Recipients"
    assert ws.max_row >= 3  # Header + 2 data rows


def test_bounced_mails_and_strict_historical_bounce_rate_invariance(client: TestClient, db: Session):
    """CRITICAL TEST:
    Verifies that deleting bounced-mail records (single or delete-all)
    removes records from the visible listing table, but NEVER alters
    the campaign's historical bounce rate or historical bounce count!
    """
    uid = uuid.uuid4().hex[:8]
    now = _utc_now()
    campaign = Campaign(
        name=f"Immunology Global Outreach {uid}",
        status="sent",
        sent_at=now,
        total_recipients=5,
        sent_count=5,
        delivered_count=2,
        opened_count=1,
        clicked_count=0,
        bounced_count=3,
    )
    db.add(campaign)
    db.flush()

    # Create 3 bounced recipients and 2 delivered recipients (5 total attempted)
    lead_b1 = Lead(name="Prof. Hard Bounce", email=f"hard_{uid}@invalid.org", country="Germany", profession="Scientist")
    lead_b2 = Lead(name="Dr. Soft Bounce", email=f"soft_{uid}@fullmailbox.com", country="France", profession="Doctor")
    lead_b3 = Lead(name="Dr. Unknown Bounce", email=f"unknown_{uid}@mystery.edu", country="Germany", profession="Dean")
    lead_ok1 = Lead(name="Dr. Valid Recipient 1", email=f"valid1_{uid}@univ.edu", country="USA", profession="Professor")
    lead_ok2 = Lead(name="Dr. Valid Recipient 2", email=f"valid2_{uid}@univ.edu", country="USA", profession="Professor")
    db.add_all([lead_b1, lead_b2, lead_b3, lead_ok1, lead_ok2])
    db.flush()

    s_b1 = CampaignRecipientSnapshot(
        campaign_id=campaign.id,
        lead_id=lead_b1.id,
        name=lead_b1.name,
        normalized_email=lead_b1.email,
        status="bounced",
        dispatch_attempt_at=now,
        bounced_at=now,
        bounce_reason="550 Mailbox does not exist",
        bounce_type="hard",
    )
    s_b2 = CampaignRecipientSnapshot(
        campaign_id=campaign.id,
        lead_id=lead_b2.id,
        name=lead_b2.name,
        normalized_email=lead_b2.email,
        status="bounced",
        dispatch_attempt_at=now,
        bounced_at=now,
        bounce_reason="452 Mailbox quota exceeded",
        bounce_type="soft",
    )
    s_b3 = CampaignRecipientSnapshot(
        campaign_id=campaign.id,
        lead_id=lead_b3.id,
        name=lead_b3.name,
        normalized_email=lead_b3.email,
        status="bounced",
        dispatch_attempt_at=now,
        bounced_at=now,
        bounce_reason="Delivery timed out",
        bounce_type="unknown",
    )
    s_ok1 = CampaignRecipientSnapshot(
        campaign_id=campaign.id,
        lead_id=lead_ok1.id,
        name=lead_ok1.name,
        normalized_email=lead_ok1.email,
        status="delivered",
        dispatch_attempt_at=now,
        delivered_at=now,
    )
    s_ok2 = CampaignRecipientSnapshot(
        campaign_id=campaign.id,
        lead_id=lead_ok2.id,
        name=lead_ok2.name,
        normalized_email=lead_ok2.email,
        status="delivered",
        dispatch_attempt_at=now,
        delivered_at=now,
    )
    db.add_all([s_b1, s_b2, s_b3, s_ok1, s_ok2])
    db.commit()

    # 1. Initial State: Check Bounced Mails summary and active listing
    res_bounces = client.get(f"/api/campaign-tracking/campaigns/{campaign.id}/bounces")
    assert res_bounces.status_code == 200
    b_data = res_bounces.json()

    assert b_data["total"] == 3
    assert b_data["summary"]["total_historical_bounced"] == 3
    assert b_data["summary"]["hard_bounces"] == 1
    assert b_data["summary"]["soft_bounces"] == 1
    assert b_data["summary"]["unknown_bounces"] == 1
    assert b_data["summary"]["active_listing_count"] == 3
    # 3 bounces out of 5 attempted = 60.0%
    assert b_data["summary"]["historical_bounce_rate"] == 60.0

    # 2. Individual Delete: Delete one bounce record (s_b1)
    res_del_single = client.delete(f"/api/campaign-tracking/campaigns/{campaign.id}/bounces/{s_b1.id}")
    assert res_del_single.status_code == 200
    del_data = res_del_single.json()
    assert del_data["success"] is True
    # Invariant: Historical bounce count and rate are unchanged!
    assert del_data["historical_bounced_count"] == 3
    assert del_data["historical_bounce_rate"] == 60.0

    # 3. Query Bounced Mails after individual delete:
    # Active listing must now show 2 records, BUT historical summary must remain 3 bounces / 60.0%!
    res_bounces_after_single = client.get(f"/api/campaign-tracking/campaigns/{campaign.id}/bounces")
    assert res_bounces_after_single.status_code == 200
    b_data2 = res_bounces_after_single.json()
    assert b_data2["total"] == 2  # Active listing reduced to 2
    assert b_data2["summary"]["active_listing_count"] == 2
    # Historical invariants remain intact:
    assert b_data2["summary"]["total_historical_bounced"] == 3
    assert b_data2["summary"]["historical_bounce_rate"] == 60.0
    assert b_data2["summary"]["hard_bounces"] == 1
    assert b_data2["summary"]["soft_bounces"] == 1

    # 4. Delete All: Delete all remaining active bounce records
    res_del_all = client.delete(f"/api/campaign-tracking/campaigns/{campaign.id}/bounces")
    assert res_del_all.status_code == 200
    del_all_data = res_del_all.json()
    assert del_all_data["success"] is True
    assert del_all_data["deleted_count"] == 2
    # Invariant: Historical bounce count and rate STILL completely unchanged!
    assert del_all_data["historical_bounced_count"] == 3
    assert del_all_data["historical_bounce_rate"] == 60.0

    # 5. Query Bounced Mails after Delete All:
    # Active listing is now completely empty (0), BUT summary remains 3 bounces / 60.0%!
    res_bounces_after_all = client.get(f"/api/campaign-tracking/campaigns/{campaign.id}/bounces")
    assert res_bounces_after_all.status_code == 200
    b_data3 = res_bounces_after_all.json()
    assert b_data3["total"] == 0  # Table is empty
    assert len(b_data3["items"]) == 0
    assert b_data3["summary"]["active_listing_count"] == 0
    # Strict historical invariant check:
    assert b_data3["summary"]["total_historical_bounced"] == 3
    assert b_data3["summary"]["historical_bounce_rate"] == 60.0
    assert b_data3["summary"]["hard_bounces"] == 1
    assert b_data3["summary"]["soft_bounces"] == 1

    # 6. Campaign Detail Summary page reload verification:
    res_summary_reload = client.get(f"/api/campaign-tracking/campaigns/{campaign.id}/summary")
    assert res_summary_reload.status_code == 200
    sum_reload = res_summary_reload.json()
    assert sum_reload["metrics"]["historical_bounced_count"] == 3
    assert sum_reload["metrics"]["historical_bounce_rate"] == 60.0
    assert sum_reload["metrics"]["active_bounce_listing_count"] == 0

    # 7. Excel export when listing empty
    res_b_excel = client.get(f"/api/campaign-tracking/campaigns/{campaign.id}/bounces/export")
    assert res_b_excel.status_code == 200
    wb_b = openpyxl.load_workbook(io.BytesIO(res_b_excel.content))
    assert wb_b.active.title == "Bounced Mails"
