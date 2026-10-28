from typing import Any
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.contact_list import ContactList, ContactListMember
from app.models.lead import Lead
from app.repositories.audience_repository import (
    create_contact_list,
    delete_contact_list,
    get_contact_filters,
    get_contact_list_by_id,
    get_contact_list_by_name,
    get_contact_lists_paginated,
    get_eligible_contacts_paginated,
    get_list_members_paginated,
    touch_contact_list,
    update_contact_list,
)
from app.schemas.audience import (
    ContactFiltersResponse,
    ContactListCreate,
    ContactListUpdate,
    MembershipBatchSummary,
)
from app.services.extraction.cleaners import (
    clean_person_name,
    is_generic_contact_email,
    is_valid_email,
)


def fetch_contacts(
    db: Session,
    page: int = 1,
    page_size: int = 20,
    search: str | None = None,
    country: str | None = None,
    university: str | None = None,
    department: str | None = None,
    profession: str | None = None,
    job_id: int | None = None,
    sort_by: str = "created_at",
    sort_order: str = "desc",
) -> tuple[list[Lead], int, int]:
    """Retrieve filtered and paginated eligible contacts."""
    return get_eligible_contacts_paginated(
        db=db,
        page=page,
        page_size=page_size,
        search=search,
        country=country,
        university=university,
        department=department,
        profession=profession,
        job_id=job_id,
        sort_by=sort_by,
        sort_order=sort_order,
    )


def fetch_contact_filters(db: Session) -> ContactFiltersResponse:
    """Retrieve distinct filter options from eligible contacts."""
    filters = get_contact_filters(db)
    return ContactFiltersResponse(**filters)


def fetch_lists(
    db: Session,
    page: int = 1,
    page_size: int = 20,
    search: str | None = None,
    filter_type: str | None = None,
    updated_within: str | None = None,
    sort_by: str = "updated_at",
    sort_order: str = "desc",
) -> tuple[list[dict[str, Any]], int, int]:
    """Retrieve filtered and paginated saved contact lists."""
    return get_contact_lists_paginated(
        db=db,
        page=page,
        page_size=page_size,
        search=search,
        filter_type=filter_type,
        updated_within=updated_within,
        sort_by=sort_by,
        sort_order=sort_order,
    )


def fetch_list_by_id(db: Session, list_id: int) -> dict[str, Any] | None:
    """Retrieve details and contact count for a single contact list."""
    return get_contact_list_by_id(db, list_id)


def create_new_list(
    db: Session,
    data: ContactListCreate,
) -> tuple[dict[str, Any], MembershipBatchSummary | None]:
    """Create a new contact list, checking for duplicate names and optionally adding initial contacts
    from existing leads or custom contact inputs.

    Multi-step creation is performed within the caller's transaction.
    """
    cleaned_name = data.name.strip()
    existing = get_contact_list_by_name(db, cleaned_name)
    if existing:
        raise ValueError(f"A list named '{cleaned_name}' already exists.")

    new_list = create_contact_list(
        db=db,
        name=cleaned_name,
        description=data.description,
        list_type=getattr(data, "list_type", "custom") or "custom",
    )
    db.flush()

    batch_summary = None
    if data.lead_ids:
        batch_summary = add_leads_to_list(
            db=db,
            list_id=new_list.id,
            lead_ids=data.lead_ids,
        )

    # Process custom contacts if provided
    if getattr(data, "contacts", None):
        custom_added = 0
        custom_skipped = 0
        seen_in_batch = set()
        for c in data.contacts:
            norm_email = str(c.email).strip().lower()
            if not norm_email or norm_email in seen_in_batch:
                custom_skipped += 1
                continue
            seen_in_batch.add(norm_email)

            # Check if Lead already exists to reuse record and preserve lead data
            lead = db.scalar(select(Lead).where(Lead.email == norm_email))
            if not lead:
                lead = Lead(
                    name=c.name.strip(),
                    email=norm_email,
                    profession=c.profession.strip() if c.profession else None,
                    university=c.university.strip() if c.university else None,
                    country=c.country.strip() if c.country else None,
                    qualification_status="qualified",
                    confidence_score=1.0,
                )
                db.add(lead)
                db.flush()
            else:
                # Update missing fields non-destructively
                if c.profession and not lead.profession:
                    lead.profession = c.profession.strip()
                if c.university and not lead.university:
                    lead.university = c.university.strip()
                if c.country and not lead.country:
                    lead.country = c.country.strip()
                db.flush()

            # Create membership if not already present
            existing_member = db.scalar(
                select(ContactListMember).where(
                    ContactListMember.list_id == new_list.id,
                    ContactListMember.normalized_email == norm_email,
                )
            )
            if not existing_member:
                member = ContactListMember(
                    list_id=new_list.id,
                    lead_id=lead.id,
                    normalized_email=norm_email,
                )
                db.add(member)
                custom_added += 1
            else:
                custom_skipped += 1

        db.flush()
        if batch_summary:
            batch_summary.added += custom_added
            batch_summary.skipped_duplicates += custom_skipped
            batch_summary.requested += len(data.contacts)
        else:
            batch_summary = MembershipBatchSummary(
                requested=len(data.contacts),
                added=custom_added,
                skipped_duplicates=custom_skipped,
                rejected_invalid=0,
            )

    list_detail = get_contact_list_by_id(db, new_list.id)
    return list_detail, batch_summary


def parse_and_validate_contacts_file(
    file_bytes: bytes,
    filename: str,
) -> dict[str, Any]:
    """Parses uploaded CSV or XLSX file, auto-detects column mapping variations,
    validates emails and names, and partitions rows into valid, invalid, and duplicates.
    """
    import io
    import csv
    import re
    from pydantic import validate_email

    lower_name = filename.lower()
    raw_rows: list[dict[str, str]] = []

    if lower_name.endswith(".csv"):
        # Decode CSV with utf-8 fallback to latin-1
        try:
            content = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = file_bytes.decode("latin-1")

        f = io.StringIO(content)
        reader = csv.DictReader(f)
        for row in reader:
            raw_rows.append({k.strip(): str(v).strip() if v is not None else "" for k, v in row.items() if k})

    elif lower_name.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        headers = []
        for col_idx, cell in enumerate(sheet.iter_cols(1, sheet.max_column, 1, 1)):
            h = str(cell[0].value or "").strip()
            headers.append(h)

        for row in sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True):
            if not any(row):
                continue
            row_dict = {}
            for col_idx, h in enumerate(headers):
                if h:
                    val = row[col_idx]
                    row_dict[h] = str(val).strip() if val is not None else ""
            raw_rows.append(row_dict)
    else:
        raise ValueError("Unsupported file format. Please upload a .csv or .xlsx file.")

    # Column synonym mapping
    NAME_SYNONYMS = {"name", "full_name", "fullname", "contact_name", "person_name", "lead_name"}
    EMAIL_SYNONYMS = {"email", "email_address", "emailaddress", "mail", "e-mail"}
    PROF_SYNONYMS = {"profession", "job_title", "title", "designation", "role", "occupation"}
    UNIV_SYNONYMS = {"university", "institution", "college", "organization", "company", "school"}
    COUNTRY_SYNONYMS = {"country", "nation", "location"}

    def find_key(row: dict[str, str], synonyms: set[str]) -> str | None:
        for k in row.keys():
            normalized_key = re.sub(r"[^a-z0-9]", "_", k.lower()).strip("_")
            if normalized_key in synonyms:
                return k
        return None

    valid_contacts = []
    invalid_rows = []
    duplicate_rows = []
    seen_emails = set()

    for idx, row in enumerate(raw_rows, start=2): # 1-based index matching spreadsheet row
        name_k = find_key(row, NAME_SYNONYMS)
        email_k = find_key(row, EMAIL_SYNONYMS)
        prof_k = find_key(row, PROF_SYNONYMS)
        univ_k = find_key(row, UNIV_SYNONYMS)
        country_k = find_key(row, COUNTRY_SYNONYMS)

        name_val = row.get(name_k, "").strip() if name_k else ""
        email_val = row.get(email_k, "").strip() if email_k else ""
        prof_val = row.get(prof_k, "").strip() if prof_k else ""
        univ_val = row.get(univ_k, "").strip() if univ_k else ""
        country_val = row.get(country_k, "").strip() if country_k else ""

        # Validation checks
        if not name_val:
            invalid_rows.append({
                "row_number": idx,
                "raw_data": row,
                "reason": "Name is missing or empty.",
            })
            continue

        if not email_val:
            invalid_rows.append({
                "row_number": idx,
                "raw_data": row,
                "reason": "Email address is missing or empty.",
            })
            continue

        norm_email = email_val.lower()
        if not re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", norm_email):
            invalid_rows.append({
                "row_number": idx,
                "raw_data": row,
                "reason": f"Invalid email format: '{email_val}'.",
            })
            continue

        if norm_email in seen_emails:
            duplicate_rows.append({
                "row_number": idx,
                "raw_data": row,
                "reason": f"Duplicate email '{norm_email}' in file.",
            })
            continue

        seen_emails.add(norm_email)
        valid_contacts.append({
            "name": name_val,
            "email": norm_email,
            "profession": prof_val or None,
            "university": univ_val or None,
            "country": country_val or None,
        })

    return {
        "total_rows": len(raw_rows),
        "valid_count": len(valid_contacts),
        "invalid_count": len(invalid_rows),
        "duplicate_count": len(duplicate_rows),
        "valid_contacts": valid_contacts,
        "invalid_rows": invalid_rows,
        "duplicate_rows": duplicate_rows,
    }



def update_existing_list(
    db: Session,
    list_id: int,
    data: ContactListUpdate,
) -> dict[str, Any] | None:
    """Update contact list metadata."""
    if data.name is not None:
        cleaned_name = data.name.strip()
        existing = get_contact_list_by_name(db, cleaned_name)
        if existing and existing.id != list_id:
            raise ValueError(f"A list named '{cleaned_name}' already exists.")

    updated = update_contact_list(
        db=db,
        list_id=list_id,
        name=data.name,
        description=data.description,
    )
    if not updated:
        return None

    db.flush()
    return get_contact_list_by_id(db, list_id)


def remove_list(db: Session, list_id: int) -> bool:
    """Delete a contact list and its member references without deleting leads."""
    return delete_contact_list(db, list_id)


def fetch_list_members(
    db: Session,
    list_id: int,
    page: int = 1,
    page_size: int = 20,
    search: str | None = None,
    sort_by: str = "added_at",
    sort_order: str = "desc",
) -> tuple[ContactList, list[dict[str, Any]], int, int] | None:
    """Fetch paginated members of a specific contact list."""
    contact_list = db.get(ContactList, list_id)
    if not contact_list:
        return None

    items, total, total_pages = get_list_members_paginated(
        db=db,
        list_id=list_id,
        page=page,
        page_size=page_size,
        search=search,
        sort_by=sort_by,
        sort_order=sort_order,
    )
    return contact_list, items, total, total_pages


def is_lead_eligible_contact(lead: Lead) -> bool:
    """Authoritative eligibility check for contact membership:

    - qualification_status == 'qualified'
    - valid RFC email address
    - non-generic role email address
    - non-blank person name
    """
    if lead.qualification_status != "qualified":
        return False
    if not lead.email or not lead.email.strip():
        return False
    clean_email = lead.email.strip().lower()
    if not is_valid_email(clean_email):
        return False
    if is_generic_contact_email(clean_email):
        return False
    if not lead.name or not lead.name.strip() or clean_person_name(lead.name) is None:
        return False
    return True


def add_leads_to_list(
    db: Session,
    list_id: int,
    lead_ids: list[int],
) -> MembershipBatchSummary:
    """Add selected leads to a contact list with rigorous duplicate prevention and eligibility validation.

    Prevents duplicate lead IDs and duplicate normalized emails within the same list.
    """
    contact_list = db.get(ContactList, list_id)
    if not contact_list:
        raise ValueError(f"Contact list with ID {list_id} does not exist.")

    # Fetch existing members for this list to identify duplicates
    existing_members_stmt = select(
        ContactListMember.lead_id,
        ContactListMember.normalized_email,
    ).where(ContactListMember.list_id == list_id)
    existing_rows = db.execute(existing_members_stmt).all()

    existing_lead_ids = {row[0] for row in existing_rows}
    existing_emails = {row[1].strip().lower() for row in existing_rows}

    # Fetch candidate leads in one batch query
    unique_requested_ids = list(dict.fromkeys(lead_ids))  # preserve order without duplicates
    leads_stmt = select(Lead).where(Lead.id.in_(unique_requested_ids))
    found_leads = {l.id: l for l in db.scalars(leads_stmt).all()}

    added_lead_ids: list[int] = []
    skipped_lead_ids: list[int] = []
    invalid_lead_ids: list[int] = []

    for lid in unique_requested_ids:
        lead = found_leads.get(lid)
        if not lead:
            invalid_lead_ids.append(lid)
            continue

        # Check authoritative eligibility rules
        if not is_lead_eligible_contact(lead):
            invalid_lead_ids.append(lid)
            continue

        normalized_email = lead.email.strip().lower()

        # Check duplicate lead ID or duplicate normalized email
        if lid in existing_lead_ids or normalized_email in existing_emails:
            skipped_lead_ids.append(lid)
            continue

        # Valid, eligible, non-duplicate member
        member = ContactListMember(
            list_id=list_id,
            lead_id=lead.id,
            normalized_email=normalized_email,
        )
        db.add(member)

        # Register in active sets to prevent duplicates within the batch
        existing_lead_ids.add(lid)
        existing_emails.add(normalized_email)
        added_lead_ids.append(lid)

    if added_lead_ids:
        touch_contact_list(db, list_id)
        db.flush()

    return MembershipBatchSummary(
        requested=len(unique_requested_ids),
        added=len(added_lead_ids),
        skipped_duplicates=len(skipped_lead_ids),
        rejected_invalid=len(invalid_lead_ids),
        added_lead_ids=added_lead_ids,
        skipped_lead_ids=skipped_lead_ids,
        invalid_lead_ids=invalid_lead_ids,
    )


def remove_leads_from_list(
    db: Session,
    list_id: int,
    lead_ids: list[int],
) -> list[int]:
    """Remove memberships of specified lead IDs from a contact list.

    Source lead records in the leads table are never modified or deleted.
    """
    contact_list = db.get(ContactList, list_id)
    if not contact_list:
        raise ValueError(f"Contact list with ID {list_id} does not exist.")

    members_stmt = (
        select(ContactListMember)
        .where(ContactListMember.list_id == list_id)
        .where(ContactListMember.lead_id.in_(lead_ids))
    )
    members_to_delete = list(db.scalars(members_stmt).all())
    removed_ids = [m.lead_id for m in members_to_delete]

    for member in members_to_delete:
        db.delete(member)

    if removed_ids:
        touch_contact_list(db, list_id)
        db.flush()

    return removed_ids
