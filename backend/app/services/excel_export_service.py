import io
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.lead import Lead
from app.models.search_result import SearchResult
from app.repositories.lead_repository import get_leads_paginated


def sanitize_cell_value(val: Any) -> Any:
    """Neutralize formula injection by prefixing risky leading characters with a single quote."""
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        return val
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")

    raw = str(val)
    if raw and raw[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{raw}"
    text = raw.strip()
    if text and text[0] in ("=", "+", "-", "@"):
        return f"'{text}"
    return text


def generate_leads_excel_bytes(leads: list[Lead]) -> bytes:
    """Generate a genuine Microsoft Excel (.xlsx) workbook for a list of leads using openpyxl.

    Applies:
      - Formula injection protection
      - Professional executive styling (navy header, white text, zebra rows)
      - Responsive column auto-width calculation
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sitech Verified Leads"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Name",
        "Email",
        "Designation / Role",
        "Specialization / Profession",
        "University / Institution",
        "Country",
        "LinkedIn URL",
        "Institutional Webpage / Lab URL",
        "Lead ID",
        "Department",
        "Phone Number",
        "Source URL",
        "Qualification Status",
        "Confidence Score",
        "Qualification Reason",
        "Search Job ID",
        "Created At (UTC)",
    ]

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border_side = Side(style="thin", color="E2E8F0")
    thin_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    regular_align = Alignment(vertical="center")
    center_align = Alignment(horizontal="center", vertical="center")

    for row_idx, lead in enumerate(leads, start=2):
        ws.row_dimensions[row_idx].height = 22
        is_even_row = (row_idx % 2 == 0)

        conf_str = ""
        if lead.confidence_score is not None:
            pct = round(lead.confidence_score * 100 if lead.confidence_score <= 1.0 else lead.confidence_score)
            conf_str = f"{pct}%"

        created_str = lead.created_at.strftime("%Y-%m-%d %H:%M:%S") if lead.created_at else ""

        row_values = [
            sanitize_cell_value(lead.name),
            sanitize_cell_value(lead.email),
            sanitize_cell_value(lead.profession),
            sanitize_cell_value(lead.research_specialization),
            sanitize_cell_value(lead.university),
            sanitize_cell_value(lead.country),
            sanitize_cell_value(lead.linkedin_url),
            sanitize_cell_value(lead.website),
            lead.id,
            sanitize_cell_value(lead.department),
            sanitize_cell_value(lead.phone_number),
            sanitize_cell_value(lead.source_url),
            sanitize_cell_value(lead.qualification_status.title() if lead.qualification_status else ""),
            conf_str,
            sanitize_cell_value(lead.qualification_reason),
            lead.search_job_id or "",
            created_str,
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if is_even_row:
                cell.fill = zebra_fill

            # Alignment rules
            if col_idx in (6, 9, 13, 14, 16, 17):
                cell.alignment = center_align
            else:
                cell.alignment = regular_align

    # Auto-adjust column widths based on maximum content length
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            cell_len = len(str(cell.value or ""))
            if cell_len > max_len:
                max_len = cell_len
        ws.column_dimensions[col_letter].width = max(12, min(50, max_len + 3))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_leads_to_excel(
    db: Session,
    status: str | None = "qualified",
    job_id: int | None = None,
) -> bytes:
    """Query leads with filtering and return an Excel (.xlsx) binary."""
    leads, _, _ = get_leads_paginated(
        db=db,
        page=1,
        page_size=20000,
        status=status,
        job_id=job_id,
    )
    if status == "qualified":
        from app.services.extraction.cleaners import is_generic_contact_email, is_valid_email
        leads = [
            lead for lead in leads
            if lead.email and is_valid_email(lead.email) and not is_generic_contact_email(lead.email)
        ]
    return generate_leads_excel_bytes(leads)


def generate_search_results_excel_bytes(results: list[SearchResult]) -> bytes:
    """Generate a genuine Microsoft Excel (.xlsx) workbook for search/extraction results using openpyxl.

    Applies:
      - Formula injection protection
      - Professional monochrome styling
      - Responsive column auto-width calculation
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stored Extraction Results"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Result ID",
        "Search Job ID",
        "Target Webpage URL",
        "Webpage Title",
        "Status",
        "Retry Count",
        "Extraction Details",
        "Associated Lead ID",
        "Processed At (UTC)",
        "Created At (UTC)",
    ]

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border_side = Side(style="thin", color="E4E4E7")
    thin_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10, color="18181B")
    regular_align = Alignment(vertical="center")
    center_align = Alignment(horizontal="center", vertical="center")

    for row_idx, r in enumerate(results, start=2):
        ws.row_dimensions[row_idx].height = 22
        is_even_row = (row_idx % 2 == 0)

        processed_str = r.processed_at.strftime("%Y-%m-%d %H:%M:%S") if r.processed_at else ""
        created_str = r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else ""
        details_str = r.error_message or ("Contact details verified & saved" if r.status == "processed" else "Awaiting parser")

        row_values = [
            r.id,
            r.search_job_id,
            sanitize_cell_value(r.result_url),
            sanitize_cell_value(r.result_title),
            sanitize_cell_value(r.status.title() if r.status else ""),
            r.retry_count or 0,
            sanitize_cell_value(details_str),
            r.lead_id or "",
            processed_str,
            created_str,
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if is_even_row:
                cell.fill = zebra_fill

            # Alignment rules
            if col_idx in (1, 2, 5, 6, 8, 9, 10):
                cell.alignment = center_align
            else:
                cell.alignment = regular_align

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            cell_len = len(str(cell.value or ""))
            if cell_len > max_len:
                max_len = cell_len
        ws.column_dimensions[col_letter].width = max(12, min(60, max_len + 3))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_search_results_to_excel(
    db: Session,
    job_id: int,
) -> bytes:
    """Query all search results for a specific job and return an Excel (.xlsx) binary."""
    from app.repositories.search_result_repository import get_all_search_results_for_job
    results = get_all_search_results_for_job(db=db, search_job_id=job_id)
    return generate_search_results_excel_bytes(results)

